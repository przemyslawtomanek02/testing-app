import os
from datetime import timezone
from io import BytesIO
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from docx import Document
from fpdf import FPDF, XPos, YPos
from sqlalchemy.orm import selectinload
from starlette import status

from fastapi_app.database.database import get_session
from fastapi_app.database.dbmodels import UserActivity, TestInstance, Result, Question
from fastapi_app.utils.dependencies import is_admin

app_dir = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
FONT_PATH = os.path.join(app_dir, "fonts", "arial.ttf")

router = APIRouter(tags=["Admin - Download Results"])

@router.post('/api/admin/download_result')
async def download_results_file(
    request: Request,
    data: dict,
    db: AsyncSession = Depends(get_session),
    _: str = Depends(is_admin)
):
    """
    Pobiera wyniki testu lub pojedynczego użytkownika i generuje plik PDF lub DOCX.
    """
    item_id = data.get("id")
    file_type = data.get("file_type", "pdf")
    entity_type = data.get("entity_type", "instance")

    if not item_id:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Brak ID")

    rows = []

    if entity_type == "user":
        stmt = (
            select(UserActivity)
            .options(
                selectinload(UserActivity.results)
                .selectinload(Result.question)
                .selectinload(Question.answers),
                selectinload(UserActivity.test_instance)
            )
            .where(UserActivity.activity_id == item_id)
        )
        result = await db.execute(stmt)
        activity = result.scalar_one_or_none()

        if not activity:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Nie znaleziono danych")

        for result_obj in activity.results:
            question = result_obj.question
            selected_answer = ""
            if question and question.answers:
                correct_answers = [a.answer_text for a in question.answers if a.is_correct and a.answer_text]
                selected_answer = ", ".join(correct_answers)

            rows.append({
                "instance_name": activity.test_instance.instance_name if activity.test_instance else "",
                "user_name": activity.user_name,
                "user_surname": activity.user_surname,
                "user_index": activity.user_index,
                "ua_timestamp": activity.timestamp,
                "test_time": activity.test_instance.test_time if activity.test_instance else 0,
                "question_number": activity.question_number,
                "score": activity.score,
                "max_instance_score": activity.test_instance.num_questions if activity.test_instance else 0,
                "question_text": question.question_text if question else "",
                "user_response": result_obj.user_response,
                "r_timestamp": result_obj.timestamp,
                "selected_answer": selected_answer
            })

    elif entity_type == "instance":
        stmt = (
            select(TestInstance)
            .options(selectinload(TestInstance.activities))
            .where(TestInstance.instance_id == item_id)
        )
        result = await db.execute(stmt)
        instance = result.scalar_one_or_none()

        if not instance:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Nie znaleziono danych")

        for activity in instance.activities:
            rows.append({
                "instance_name": instance.instance_name,
                "user_name": activity.user_name,
                "user_surname": activity.user_surname,
                "user_index": activity.user_index,
                "ua_timestamp": activity.timestamp,
                "test_time": instance.test_time,
                "score": activity.score,
                "max_instance_score": instance.num_questions
            })

    if not rows:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Brak danych do wygenerowania pliku")

    file_stream, filename = generate_file(rows, entity_type, file_type)

    filename_header = f"filename*=UTF-8''{quote(filename)}"

    return StreamingResponse(
        file_stream,
        media_type=f"application/{file_type}",
        headers={"Content-Disposition": f'attachment; {filename_header}'}
    )


async def fetch_data(db: AsyncSession, entity_type: str, item_id: str):
    """Pobiera dane z bazy asynchronicznie i zwraca listę słowników."""

    if entity_type == "user":
        stmt = select(UserActivity).where(UserActivity.activity_id == item_id)
        result = await db.execute(stmt)
        activity = result.scalar_one_or_none()
        if not activity:
            return []

        rows = []
        for result_obj in activity.results:
            question = result_obj.question
            selected_answer = ""
            if question and question.answers:
                correct_answers = [a.answer_text for a in question.answers if a.is_correct and a.answer_text]
                selected_answer = ", ".join(correct_answers)
            rows.append({
                "instance_name": activity.test_instance.instance_name if activity.test_instance else "",
                "user_name": activity.user_name,
                "user_surname": activity.user_surname,
                "user_index": activity.user_index,
                "ua_timestamp": activity.timestamp,
                "test_time": activity.test_instance.test_time if activity.test_instance else 0,
                "question_number": activity.question_number,
                "score": activity.score,
                "max_instance_score": activity.test_instance.num_questions if activity.test_instance else 0,
                "question_text": question.question_text if question else "",
                "user_response": result_obj.user_response,
                "r_timestamp": result_obj.timestamp,
                "selected_answer": selected_answer
            })
        return rows

    elif entity_type == "instance":
        stmt = select(TestInstance).where(TestInstance.instance_id == item_id)
        result = await db.execute(stmt)
        instance = result.scalar_one_or_none()
        if not instance:
            return []

        rows = []
        for activity in instance.activities:
            rows.append({
                "instance_name": instance.instance_name,
                "user_name": activity.user_name,
                "user_surname": activity.user_surname,
                "user_index": activity.user_index,
                "ua_timestamp": activity.timestamp,
                "test_time": instance.test_time,
                "score": activity.score,
                "max_instance_score": instance.num_questions
            })
        return rows
    else:
        return []


def generate_file(rows, entity_type, file_type):
    first_row = rows[0]

    if entity_type == "user":
        filename = f"{first_row['user_surname']}_{first_row['user_name']}_{first_row['user_index']}_{first_row['instance_name']}_results.{file_type}"
    else:
        filename = f"{first_row['instance_name']}_results.{file_type}"

    file_stream = BytesIO()

    if file_type == "pdf":
        if entity_type == "user":
            generate_user_pdf(rows, file_stream)
        else:
            generate_instance_pdf(rows, file_stream)
    elif file_type == "docx":
        if entity_type == "user":
            generate_user_docx(rows, file_stream)
        else:
            generate_instance_docx(rows, file_stream)

    file_stream.seek(0)
    return file_stream, filename

def generate_user_pdf(rows, file_stream):
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("Arial", "", FONT_PATH)
    pdf.set_font("Arial", size=12)

    # Nagłówek raportu
    pdf.write(10, f"Raport użytkownika - {rows[0]['user_name']} {rows[0]['user_surname']} ({rows[0]['user_index']})\n")
    pdf.write(10, f"Test: {rows[0]['instance_name']}\n")
    pdf.write(10, f"Czas testu: {rows[0]['test_time']} minut\n")
    pdf.write(10, f"Wynik: {rows[0]['score']} / {rows[0]['max_instance_score']}\n")
    pdf.write(10, f"Użytkownik odpowiedział na: {rows[0]['question_number']} pytań\n\n")

    pdf.set_auto_page_break(auto=True, margin=10)

    for index, row in enumerate(rows, start=1):
        text = (f"{index}. Pytanie: {row['question_text']}\n"
                f"Odpowiedź: {row['user_response']}\n"
                f"Poprawna odpowiedź: {row['selected_answer']}\n")
        pdf.multi_cell(0, 7, text)

    pdf.output(file_stream)


def generate_instance_pdf(rows, file_stream):
    """Generuje PDF dla instancji testu."""

    rows_sorted = sorted(rows, key=lambda x: x['user_surname'].lower())

    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("Arial", "", FONT_PATH)
    pdf.set_font("Arial", size=12)

    pdf.cell(200, 10, f"Raport testu - {rows[0]['instance_name']}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(200, 10, f"Liczba przystępujących - {len(rows)}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(5)

    col_widths = [80, 40, 60]
    pdf.set_font("Arial", size=10)

    pdf.cell(col_widths[0], 10, "Użytkownik", border=1, align='C')
    pdf.cell(col_widths[1], 10, "Wynik", border=1, align='C')
    pdf.cell(col_widths[2], 10, "Start testu", border=1, align='C')
    pdf.ln()

    # Dane użytkowników
    for row in rows_sorted:
        pdf.cell(col_widths[0], 10, f"{row['user_surname']} {row['user_name']} ({row['user_index']})", border=1)
        pdf.cell(col_widths[1], 10, f"{row['score']} / {row['max_instance_score']}", border=1, align='C')
        ua_timestamp = row['ua_timestamp']
        if ua_timestamp.tzinfo is None:
            ua_timestamp = ua_timestamp.replace(tzinfo=timezone.utc)

        local_time = ua_timestamp.astimezone(ZoneInfo("Europe/Warsaw"))
        formatted = local_time.strftime("%d-%m-%Y %H:%M")
        pdf.cell(col_widths[2], 10, formatted, border=1, align='C')
        pdf.ln()

    pdf.output(file_stream)


def generate_user_docx(rows, file_stream):
    """Generuje DOCX dla użytkownika."""
    doc = Document()

    doc.add_heading(f"Raport użytkownika - {rows[0]['user_name']} {rows[0]['user_surname']} ({rows[0]['user_index']})",
                    level=1)

    doc.add_paragraph(f"Test: {rows[0]['instance_name']}")
    doc.add_paragraph(f"Czas testu: {rows[0]['test_time']} minut")
    doc.add_paragraph(f"Wynik: {rows[0]['score']} / {rows[0]['max_instance_score']}")
    doc.add_paragraph(f"Użytkownik odpowiedział na: {rows[0]['question_number']} pytań")
    doc.add_paragraph("-----------------------------------------")

    for index, row in enumerate(rows, start=1):
        doc.add_paragraph(f"{index}. Pytanie: {row['question_text']}")
        doc.add_paragraph(f"Odpowiedź: {row['user_response']}")
        doc.add_paragraph(f"Poprawna odpowiedź: {row['selected_answer']}")
        doc.add_paragraph("-----------------------------------------")

    doc.save(file_stream)


def generate_instance_docx(rows, file_stream):
    """Generuje DOCX dla instancji testu."""

    rows_sorted = sorted(rows, key=lambda x: x['user_surname'].lower())

    doc = Document()

    doc.add_heading(f"Raport testu - {rows[0]['instance_name']}", level=1)

    doc.add_paragraph(f"Liczba przystępujących - {len(rows)}")
    doc.add_paragraph("-----------------------------------------")

    for row in rows_sorted:
        ua_timestamp = row['ua_timestamp']
        if ua_timestamp.tzinfo is None:
            ua_timestamp = ua_timestamp.replace(tzinfo=timezone.utc)

        local_time = ua_timestamp.astimezone(ZoneInfo("Europe/Warsaw"))
        formatted = local_time.strftime("%d-%m-%Y %H:%M")
        doc.add_paragraph(f"{row['user_surname']} {row['user_name']} (Indeks: {row['user_index']})\n"
                          f"Wynik: {row['score']} / {row['max_instance_score']}     Start testu: {formatted}\n"
                          f"-----------------------------------------\n")

    doc.save(file_stream)
